#flask web server - handles http, generates proposals, etc
from flask import Flask, request, jsonify, send_file
from docx import Document
import os

app = Flask(__name__)

# Route to generate the proposal
@app.route('/generate-proposal', methods=['POST'])
def generate_proposal():
    data = request.json  # Expect client data from React app
    client_name = data.get('name', 'Unknown Client')
    project_details = data.get('details', 'No details provided')

    # Load the Word template
    template_path = 'Proposal Inspek.docx'
    doc = Document(template_path)

    # Replace placeholders in the document
    for paragraph in doc.paragraphs:
        if "{{client_name}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{client_name}}", client_name)
        if "{{project_details}}" in paragraph.text:
            paragraph.text = paragraph.text.replace("{{project_details}}", project_details)

    # Save the customized document
    output_path = os.path.join('output', f'{client_name}_Proposal.docx')
    doc.save(output_path)

    return send_file(output_path, as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True)
