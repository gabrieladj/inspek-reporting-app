import os
from docx import Document
from datetime import datetime
from dotenv import load_dotenv
import sys
import pymongo
from bson import ObjectId  # Import ObjectId from bson
import argparse
from io import BytesIO
# from docxtpl import DocxTemplate

# Load environment variables
load_dotenv()

# Path to your proposal template
TEMPLATE_PATH = os.getenv('TEMPLATE_PATH')
if TEMPLATE_PATH is None:
    print("Error: TEMPLATE_PATH is not set in the environment variables!")
    sys.exit(1)


# Normalize path separator for cross-platform compatibility
TEMPLATE_FILE = os.path.join(TEMPLATE_PATH, 'proposal_template.docx')
print(f"TEMPLATE_PATH: {TEMPLATE_PATH}")

# Connect to MongoDB (adjust connection details as needed)
client = pymongo.MongoClient(os.getenv("MONGODB_URI"))
db = client.get_database()  # Use the correct database name
client_collection = db["clients"]  
report_collection = db["reports"]  
print(f"Connected to DB...")

# Parse arguments
parser = argparse.ArgumentParser(description="Generate a proposal document")
parser.add_argument("client_id", help="Client ID")
parser.add_argument("report_id", help="Report ID")
parser.add_argument("clientName", help="Client name")
parser.add_argument("mailingAddress", help="Client mailing address")
parser.add_argument("propertyName", help="Property name")
parser.add_argument("propertyAddress", help="Property address")
parser.add_argument("officeSpacePercentage", help="Percentage of office space")
parser.add_argument("warehouseSpacePercentage", help="Percentage of warehouse space")
parser.add_argument("retailSpacePercentage", help="Percentage of retail space")
parser.add_argument("otherSpacePercentage", help="Percentage of other space")
parser.add_argument("propertyRepresentativeName", help="Property representative name")
parser.add_argument("manufacturingSpacePercentage", help="Percentage of manufacturing space")
parser.add_argument("roleOrRelationship", help="Role or relationship of the client")
parser.add_argument("propertyType", help="Type of property")
parser.add_argument("totalBuildingSqFt", help="Total building sq. ft.")
parser.add_argument("typeOfInspection", help="Inspection type")
args = parser.parse_args()

print("Arguments passed to the script:", sys.argv)  # Check what data is received

client_id = args.client_id
report_id = args.report_id

# Validate IDs
if not ObjectId.is_valid(client_id):
    print(f"Error: Invalid client_id: {client_id}")
    sys.exit(1)
if not ObjectId.is_valid(report_id):
    print(f"Error: Invalid report_id: {report_id}")
    sys.exit(1)

def generate_proposal(client_name, mailing_address, property_name,
                      property_address, office_space_percentage, warehouse_space_percentage,
                      retail_space_percentage, manufacturing_space_percentage, other_space_percentage, 
                      property_representative_name, role_or_relationship, type_of_property, total_building_sq_ft, inspection_type):

    # Convert clientId and reportId to ObjectId
    client_id = ObjectId(args.client_id)
    report_id = ObjectId(args.report_id)

    # Fetch data from MongoDB
    print(f"Querying client collection with client_id: {client_id}")
    client_data = client_collection.find_one({"_id": client_id})

    print(f"Querying report collection with report_id: {report_id}")
    report_data = report_collection.find_one({"_id": report_id})
    print("...")

    if not client_data:
        print(f"No client data found for clientId: {client_id}")
        sys.exit(1)
    if not report_data:
        print(f"No report data found for reportId: {report_id}")
        sys.exit(1)


    # Print the fetched client and report data to check if fields are correct
    # debugging lines
    print("...")
    print("FETCHED DATA:")
    print(f"Client Data: {client_data}")
    print(f"Report Data: {report_data}")
    print("END OF FETCHED DATA")
    print("...")

    # for debugging client_data and report_data
    print(f"Client Name: {client_data.get('clientName')}")
    print(f"Mailing Address: {client_data.get('mailingAddress')}")
    print(f"Property Name: {report_data.get('propertyName')}")
    print(f"Property Address: {report_data.get('propertyAddress')}")
    print(f"Office Space Percentage: {report_data.get('officeSpacePercentage')}")
    print(f"Warehouse Space Percentage: {report_data.get('warehouseSpacePercentage')}")
    print(f"Retail Space Percentage: {report_data.get('retailSpacePercentage')}")
    print(f"Other Space Percentage: {report_data.get('otherSpacePercentage')}")
    print(f"Property Representative Name: {report_data.get('propertyRepresentativeName')}")

    print(f"Manufacturing Space Percentage: {report_data.get('manufacturingSpacePercentage')}")
    print(f"Role or Relationship: {client_data.get('roleOrRelationship')}")
    print(f"Type of Property: {client_data.get('propertyType')}")
    print(f"Total Building Sq. Ft.: {report_data.get('totalBuildingSqFt')}")
    print(f"Type of Inspection: {report_data.get('typeOfInspection')}")

    # Load template
    if not os.path.exists(TEMPLATE_FILE):
        print(f"Error: Template file not found at {TEMPLATE_FILE}")
        sys.exit(1)

    # Open the template document
    doc = Document(TEMPLATE_FILE)

    # Updated placeholders extraction
    placeholders = {
        "{clientName}": client_data.get("clientName", ""),
        "{mailingAddress}": client_data.get("mailingAddress", ""),
        "{propertyName}": report_data.get("propertyInfo", {}).get("propertyName", ""),
        "{propertyAddress}": report_data.get("propertyInfo", {}).get("propertyAddress", ""),
        "{officeSpacePercentage}": str(report_data.get("buildingDetails", {}).get("officeSpacePercentage", "")),
        "{warehouseSpacePercentage}": str(report_data.get("buildingDetails", {}).get("warehouseSpacePercentage", "")),
        "{retailSpacePercentage}": str(report_data.get("buildingDetails", {}).get("retailSpacePercentage", "")),
        "{otherSpacePercentage}": str(report_data.get("buildingDetails", {}).get("otherSpacePercentage", "")),
        "{propertyRepresentativeName}": client_data.get("propertyRepresentativeName", ""),
        "{manufacturingSpacePercentage}": str(report_data.get("buildingDetails", {}).get("manufacturingSpacePercentage", "")),
        "{roleOrRelationship}": client_data.get("roleOrRelationship", ""),
        "{propertyType}": report_data.get("propertyInfo", {}).get("propertyType", ""),
        "{totalBuildingSqFt}": str(report_data.get("propertyInfo", {}).get("totalBuildingSqFt", "")),
        "{typeOfInspection}": report_data.get("inspectionScope", {}).get("typeOfInspection", "")
    }

    # Replace text in paragraphs
    for paragraph in doc.paragraphs:
        for placeholder, value in placeholders.items():
            if placeholder in paragraph.text:
                print(f"Replacing placeholder: {placeholder} with value: {value}")  # Debug log
                paragraph.text = paragraph.text.replace(placeholder, value)

    # Replace text in shapes here

    # Save the document to a BytesIO buffer
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    print(f"UPLOAD_FOLDER: {os.getenv('UPLOAD_FOLDER')}")

    upload_folder = os.getenv('UPLOAD_FOLDER', './report_scripts/uploads')
    if not os.path.exists(upload_folder):
        os.makedirs(upload_folder)  # Create the directory if it doesn't exist

    # Save the file to disk
    output_filename = f"Proposal_{args.clientName}_{args.propertyName}.docx"
    generated_file_path = os.path.join(upload_folder, output_filename)
    print(f"OUTPUT FILENAME:", (output_filename))
    print(f"Generated file path by Python: {generated_file_path}")

    print("...")

    with open(generated_file_path, 'wb') as f:
        f.write(buffer.getvalue())  # Make sure buffer content is properly written as Word doc

    # Return the GENERATED file path to the Node.js app
    sys.stdout.write(generated_file_path)

# Calling the function with arguments
generate_proposal(
    args.clientName,
    args.mailingAddress,
    args.propertyName,
    args.propertyAddress,
    args.officeSpacePercentage,
    args.warehouseSpacePercentage,
    args.retailSpacePercentage,
    args.otherSpacePercentage,
    args.propertyRepresentativeName,
    args.manufacturingSpacePercentage,
    args.roleOrRelationship,
    args.propertyType,
    args.totalBuildingSqFt,
    args.typeOfInspection
)